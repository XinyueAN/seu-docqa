from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path


@dataclass
class ParsedPage:
    page: int
    text: str
    needs_ocr: bool = False


@dataclass
class ParsedImage:
    page: int
    filename: str
    content_type: str
    data: bytes


@dataclass
class ParsedDocument:
    title: str
    kind: str
    pages: list[ParsedPage]
    images: list[ParsedImage]


def parse_document(path: Path) -> ParsedDocument:
    suffix = path.suffix.lower()
    images: list[ParsedImage] = []
    if suffix == ".pdf":
        from pypdf import PdfReader
        pages = []
        for i, page in enumerate(PdfReader(str(path)).pages, 1):
            extracted_text = (page.extract_text() or "").strip()
            pages.append(ParsedPage(i, extracted_text, needs_ocr=not bool(extracted_text)))
            for image_index, image in enumerate(page.images, 1):
                try:
                    filename = image.name or f"page-{i}-image-{image_index}.png"
                    images.append(ParsedImage(i, filename, _content_type(filename), image.data))
                except Exception:
                    continue
        kind = "PDF"
    elif suffix == ".docx":
        from docx import Document
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        document = Document(str(path))
        blocks = []
        for block in document.iter_inner_content():
            if isinstance(block, Paragraph) and block.text.strip():
                blocks.append(block.text.strip())
            elif isinstance(block, Table):
                rows = [[cell.text.strip().replace("|", "\\|").replace("\n", "<br>") for cell in row.cells] for row in block.rows]
                if rows:
                    width = max(len(row) for row in rows)
                    rows = [row + [""] * (width - len(row)) for row in rows]
                    table_md = ["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |"]
                    table_md.extend("| " + " | ".join(row) + " |" for row in rows[1:])
                    blocks.append("\n".join(table_md))
        pages = [ParsedPage(1, "\n\n".join(blocks))]
        for i, shape in enumerate(document.inline_shapes, 1):
            try:
                relation_id = shape._inline.graphic.graphicData.pic.blipFill.blip.embed
                part = document.part.related_parts[relation_id]
                filename = Path(str(part.partname)).name or f"image-{i}.png"
                images.append(ParsedImage(1, filename, part.content_type, part.blob))
            except Exception:
                continue
        kind = "Word"
    elif suffix == ".pptx":
        from pptx import Presentation
        pages = []
        for i, slide in enumerate(Presentation(str(path)).slides, 1):
            blocks = []
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.text.strip():
                    blocks.append(shape.text.strip())
                if getattr(shape, "has_table", False):
                    rows = [[cell.text.strip().replace("|", "\\|").replace("\n", "<br>") for cell in row.cells] for row in shape.table.rows]
                    if rows:
                        width = max(len(row) for row in rows)
                        rows = [row + [""] * (width - len(row)) for row in rows]
                        blocks.append("\n".join(["| " + " | ".join(rows[0]) + " |", "| " + " | ".join(["---"] * width) + " |", *["| " + " | ".join(row) + " |" for row in rows[1:]]]))
                if getattr(shape, "has_chart", False):
                    try:
                        chart = shape.chart
                        plot = chart.plots[0]
                        labels = [str(category.label) for category in plot.categories]
                        chart_rows = []
                        for series in chart.series:
                            values = [str(value) for value in series.values]
                            chart_rows.append([str(series.name), *values[:len(labels)]])
                        if labels and chart_rows:
                            blocks.append("\n".join([
                                "| 系列 | " + " | ".join(labels) + " |",
                                "| --- | " + " | ".join(["---"] * len(labels)) + " |",
                                *["| " + " | ".join(row) + " |" for row in chart_rows],
                            ]))
                    except Exception:
                        pass
                if getattr(shape, "shape_type", None) == 13:
                    try:
                        picture = shape.image
                        filename = picture.filename or f"slide-{i}-image-{len(images)+1}.{picture.ext}"
                        images.append(ParsedImage(i, filename, picture.content_type, picture.blob))
                    except Exception:
                        continue
            pages.append(ParsedPage(i, "\n\n".join(blocks)))
        kind = "PPT"
    elif suffix in {".md", ".txt"}:
        pages = [ParsedPage(1, path.read_text(encoding="utf-8", errors="replace"))]
        kind = "Markdown" if suffix == ".md" else "Text"
    else:
        raise ValueError("当前支持 PDF、DOCX、PPTX、Markdown 和 TXT 文件。")
    return ParsedDocument(title=path.name, kind=kind, pages=pages, images=images)


def _content_type(filename: str) -> str:
    ext = Path(filename).suffix.lower()
    return {".jpg":"image/jpeg", ".jpeg":"image/jpeg", ".png":"image/png", ".gif":"image/gif", ".webp":"image/webp"}.get(ext, "image/png")


def to_markdown(document: ParsedDocument) -> str:
    sections = [f"# {document.title}"]
    for page in document.pages:
        sections.append(f"\n\n## 第 {page.page} 页\n\n{page.text}" if page.text else f"\n\n## 第 {page.page} 页\n\n_未提取到文本内容。_")
    return "".join(sections)
